"""Сборка отчёта ЛР3 по семи метрикам первоначального и повторного замеров."""
from pathlib import Path
import argparse
import json

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output-dir', type=Path, default=BASE.parents[2],
                    help='Каталог готового отчёта; по умолчанию рядом с TourneyHub')
args = parser.parse_args()
args.output_dir.mkdir(parents=True, exist_ok=True)
before = json.loads((BASE / 'metrics-before.json').read_text())
after = json.loads((BASE / 'metrics-after.json').read_text())
b, a = before['summary'], after['summary']
METRICS = [
    ('Максимальная цикломатическая сложность', 'cyclomatic_max',
     'Максимум CC среди функций и методов. gocyclo учитывает базовую единицу и ветвления: if, циклы, case, select, логические && и ||.'),
    ('Максимальная когнитивная сложность', 'cognitive_max',
     'Максимум оценки gocognit среди функций и методов. Учитываются ветвления, их вложенность и последовательности логических операторов.'),
    ('Физические строки LOC', 'source_lines',
     'Все строки измеряемых файлов, включая код, комментарии и пустые строки.'),
    ('Строки кода SLOC', 'code_lines',
     'Строки с токенами Go. Комментарии и автоматически вставленные точки с запятой исключаются; многострочные литералы учитываются целиком.'),
    ('Максимальная длина функции', 'function_lines_max',
     'Максимальное число физических строк от объявления функции или метода до закрывающей скобки.'),
    ('Доля комментариев', 'comment_density_percent',
     'Число строк с комментариями, делённое на LOC и умноженное на 100 %. Строка с кодом и комментарием входит в обе категории.'),
    ('Максимальный fan out пакета', 'package_fan_out_max',
     'Максимальное число уникальных импортируемых пакетов для одного собственного пакета, включая стандартную библиотеку.'),
]
assert len(METRICS) == 7

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style._element.get_or_add_rPr().rFonts
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():
            del fonts.attrib[attr]
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + attr), 'Times New Roman')
    for border in list(style.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(7)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].font.size = Pt(15 if name == 'Heading 1' else 13)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Subtitle'].font.size = Pt(15)


def p(text='', style=None):
    return doc.add_paragraph(text, style)


def h(text):
    return doc.add_heading(text, level=1)


def page():
    doc.add_page_break()


def table(headers, rows, widths):
    result = doc.add_table(rows=1, cols=len(headers))
    result.style = 'Table Grid'
    result.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    result._tbl.tblPr.append(borders)
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    for i, name in enumerate(headers):
        result.rows[0].cells[i].text = str(name)
    result.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for row_idx, row in enumerate(result.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF2')
                cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                if len(headers) > 2 and i > 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(5)
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.bold = row_idx == 0
    p()
    return result


def value(summary, key):
    number = summary[key]
    if key == 'comment_density_percent':
        return f'{number:.2f}'.replace('.', ',') + ' %'
    return str(number)


def operation(report, name):
    matches = [item for item in report['functions'] if item['name'].split('.')[-1] == name]
    if len(matches) > 1:
        matches = [item for item in matches if item['file'].startswith('internal/service/')]
    assert len(matches) == 1, name
    return matches[0]


for text in [
    'Национальный исследовательский университет',
    'Высшая школа экономики',
    'Московский институт электроники и математики имени А Н Тихонова',
]:
    p(text).alignment = WD_ALIGN_PARAGRAPH.CENTER
p()
p()
p()
p('Лабораторная работа 3', 'Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('Метрики программного кода', 'Subtitle').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('TourneyHub Серверная часть на Go', 'Subtitle').alignment = WD_ALIGN_PARAGRAPH.CENTER
p()
p('Предварительный отчёт по backend').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('В работе измерены семь метрик серверного Go-кода TourneyHub. Первоначальные результаты использованы для выбора участков рефакторинга; повторное измерение показывает изменение сложности, объёма и связности кода.')
p('Курс Технологии разработки программного обеспечения')
p('Дата повторного измерения 26 октября 2026 года')
p('Москва 2026').alignment = WD_ALIGN_PARAGRAPH.CENTER

page()
h('1 Цель работы')
p('Оценить программный код с помощью метрик, выявить наиболее сложные участки, оптимизировать их и сравнить первоначальные и повторные результаты. Объект исследования - серверная часть проекта TourneyHub на Go. Отчёт охватывает семь метрик; результаты frontend будут добавлены в общий командный отчёт отдельно.')
h('2 Инструменты и методика')
p('Измерения выполнены на Go 1.27.2. Для цикломатической сложности использован gocyclo v0.6.0, для когнитивной - gocognit v1.2.0. Остальные показатели рассчитаны инструментом cmd/metrics с помощью стандартных go/parser, go/ast и go/scanner. Версии анализаторов закреплены в go.mod.')
p('В обоих замерах исследуется весь написанный вручную production Go-код по одинаковым правилам отбора. Исключены тесты, сгенерированные файлы, инструмент измерения, вспомогательные инструменты и каталоги сборки. Состав исходников между замерами изменился. Архивы двух состояний и SHA-256 файлов сохранены для воспроизведения расчётов.')
table(['Метрика', 'Способ расчёта'], [(label, definition) for label, _, definition in METRICS], [6, 10.5])
p('Для выбора участков анализа использованы ориентиры CC > 10, когнитивная сложность > 15 и длина функции > 60 строк. Они указывают на необходимость изучить код, но сами по себе не доказывают наличие ошибки. Анонимные функции входят в оценку сложности содержащего объявления.')

page()
h('3 Первоначальное измерение и анализ')
p('Первоначальные значения сохранены в metrics-before.json и приведены в столбце «Первоначально» итоговой таблицы. Наибольшие значения цикломатической и когнитивной сложности составили 27 и 55. Максимальная длина функции достигла 82 строк. Все три максимума относятся к DisputeResult: последовательность проверок и вложенных условий затрудняла чтение этой операции.')
p('Дополнительно изучены DecideInvitation и UpdateProfile. В них проверки входных данных, прав и состояния выполнялись вместе с применением изменений. Эти функции выбраны для выделения отдельных шагов и уменьшения вложенности. В анализе учитывалась не только итоговая оценка, но и смысл каждой ветви.')
p('Максимальный fan out равнялся 44: отдельный пакет имел большое число прямых зависимостей. Доля комментариев составляла 1,75 %. Она характеризует объём пояснений, но не позволяет оценить их полезность. Объём исходников зафиксирован как исходная точка для повторного сравнения.')
h('4 Оптимизация кода')
p('Применено выделение функций с одной задачей. В DisputeResult отдельно выражены выбор данных, проверка допустимости действия и применение изменений. В DecideInvitation разделены проверка состояния и обработка решения. В UpdateProfile вынесены проверки отдельных полей. Последовательность действий сохранена; вложенные проверки по возможности заменены ранним возвратом ошибки.')
rows = []
for name in ['DisputeResult', 'DecideInvitation', 'UpdateProfile']:
    initial, repeated = operation(before, name), operation(after, name)
    rows.append((name, f"{initial['cyclomatic']} → {repeated['cyclomatic']}", f"{initial['cognitive']} → {repeated['cognitive']}"))
table(['Функция', 'CC до и после', 'Когнитивная сложность до и после'], rows, [6.5, 4, 6])
p('Таблица показывает оценки именованных функций после выделения шагов. Выделенные вспомогательные функции также входят в повторный замер всего кода: сложность распределяется между небольшими операциями, а не исключается из измерения. Для проверки сохранения поведения после изменений выполнены автоматические проверки; make check завершился успешно.')

page()
h('5 Повторное измерение')
p('После оптимизации повторены те же расчёты с теми же версиями инструментов. Результаты сохранены в metrics-after.json. В таблице представлены только семь выбранных метрик для двух состояний кода.')
table(['Метрика', 'Первоначально', 'Повторно'], [(label, value(b, key), value(a, key)) for label, key, _ in METRICS], [8, 4.25, 4.25])
p('Максимальные значения цикломатической и когнитивной сложности снизились с 27 до 19 и с 55 до 25. Максимальная длина функции уменьшилась с 82 до 59 строк. Это показывает уменьшение наиболее высоких локальных оценок; максимумы двух замеров могут относиться к разным функциям.')
p('LOC и SLOC выросли: добавлены вспомогательные функции, явные проверки и самостоятельные объявления типов. Поэтому нельзя утверждать, что все показатели уменьшились. Рост объёма нужно рассматривать вместе с локальной сложностью. Увеличение доли комментариев также не является самостоятельным доказательством качества; снижение fan out описывает прямые зависимости пакетов и требует содержательной оценки кода.')
h('6 Выводы')
p('Выполнены все этапы ЛР3: построены семь метрик, проанализирован первоначальный результат, проведена оптимизация, выполнен повторный замер и сопоставлены показатели. Локальная сложность и максимальная длина функции снизились, общий объём кода вырос. Оставшиеся высокие оценки сложности требуют адресного анализа; дальнейшие изменения следует выбирать по понятности и сопровождаемости конкретного кода.')
p('Для воспроизведения используются source-before.tar.gz, source-after.tar.gz и cmd/metrics. Полные команды и контрольные суммы архивов приведены в README раздела ЛР3. Источники: ТРПО_Лаба3_2023.pdf, задание на странице 12; исходники TourneyHub; metrics-before.json и metrics-after.json.')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Лабораторная работа 3 Метрики программного кода TourneyHub Backend'
doc.core_properties.subject = 'Семь метрик Первоначальное и повторное измерение'
doc.core_properties.author = ''
doc.core_properties.keywords = 'TourneyHub, Go, ЛР3, метрики'
out = args.output_dir / 'TourneyHub_ЛР3_backend_предварительный_отчет.docx'
doc.save(out)
print(out)
