"""Экспорт Markdown подготовки к защите в Word рядом с репозиторием.

Зависимость: python-docx. PDF экспортируется из готового Word-документа.
"""
from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        else:
            link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            paragraph.add_run(f'{link[1]} ({link[2]})' if link else part)


def prepare():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
        for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + attr), 'Times New Roman')
        for border in list(style.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(21)
    doc.styles['Title'].paragraph_format.space_after = Pt(12)
    for level in [1, 2, 3]:
        style = doc.styles[f'Heading {level}']
        style.font.size = Pt(16 - level)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    code = doc.styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = 'Courier New'
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Cm(.25)
    code.paragraph_format.widow_control = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    return doc


def cells(line):
    return [part.strip().replace('\\|', '|') for part in re.split(r'(?<!\\)\|', line.strip().strip('|'))]


def add_table(doc, lines):
    content = [cells(lines[0])] + [cells(line) for line in lines[2:]]
    count = len(content[0])
    assert all(len(row) == count for row in content), content
    widths = {2: [5.5, 11], 3: [7, 4.75, 4.75], 4: [4.8, 3.9, 3.9, 3.9]}.get(count, [16.5 / count] * count)
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(content):
        added = table.add_row()
        added._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index == 0:
            added._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for col, (cell, text) in enumerate(zip(added.cells, row)):
            cell.width = Cm(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF2')
                cell._tc.get_or_add_tcPr().append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            inline(paragraph, text)
            if col > 0 and re.fullmatch(r'[\d\s,.%→+-]+', text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                if run.font.name != 'Courier New':
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10.5)
                if index == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_markdown(doc, source):
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('```'):
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith('```'):
                block.append(lines[index])
                index += 1
            for row, value in enumerate(block):
                paragraph = doc.add_paragraph(style='Source Code')
                paragraph.add_run(value)
                paragraph.paragraph_format.keep_with_next = len(block) <= 12 and row < len(block) - 1
                if row == len(block) - 1:
                    paragraph.paragraph_format.space_after = Pt(8)
            index += 1
            continue
        if line.startswith('|') and index + 1 < len(lines) and re.fullmatch(r'[\s|:-]+', lines[index + 1]):
            block = [lines[index], lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith('|'):
                block.append(lines[index])
                index += 1
            add_table(doc, block)
            continue
        heading = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading:
            level = len(heading[1])
            paragraph = doc.add_paragraph(style='Title' if level == 1 else f'Heading {level - 1}')
            inline(paragraph, heading[2])
            index += 1
            continue
        bullet = re.match(r'^[-*]\s+(.+)$', line)
        numbered = re.match(r'^(\d+[.)])\s+(.+)$', line)
        if bullet:
            paragraph = doc.add_paragraph(style='List Bullet')
            inline(paragraph, bullet[1])
        elif numbered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(.65)
            paragraph.paragraph_format.first_line_indent = Cm(-.65)
            inline(paragraph, numbered[1] + ' ' + numbered[2])
        else:
            text = [line]
            while index + 1 < len(lines) and lines[index + 1].strip():
                following = lines[index + 1].strip()
                if re.match(r'^(#|```|\||[-*]\s|\d+[.)]\s)', following):
                    break
                text.append(following)
                index += 1
            inline(doc.add_paragraph(), ' '.join(text))
        index += 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=BASE / 'defense-guide.md')
    parser.add_argument('--output-dir', type=Path, default=BASE.parents[2])
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    doc = prepare()
    render_markdown(doc, args.source.read_text())
    doc.core_properties.title = 'TourneyHub Подготовка к защите лабораторной работы 3'
    doc.core_properties.subject = 'Проект Backend Семь метрик Подготовка к защите'
    doc.core_properties.author = ''
    doc.core_properties.keywords = 'TourneyHub, Go, backend, ЛР3, защита, метрики'
    out = args.output_dir / 'TourneyHub_ЛР3_подготовка_к_защите.docx'
    doc.save(out)
    print(out)


if __name__ == '__main__':
    main()
